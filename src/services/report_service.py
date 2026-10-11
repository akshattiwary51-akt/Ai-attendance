"""Attendance reports for download (CSV / Excel). Numbers come from the same RLS-scoped services as the dashboards, so a
user can only ever export what they are allowed to see.

Counting rules: attended = present + late; conducted = present + late + absent; excused is neutral; only completed sessions count.
"""
from __future__ import annotations

import csv
import io
from datetime import datetime, timezone

from src.analytics import attendance_math as am
from src.services import analytics_service, dashboard_service
from src.utils.timefmt import format_local

_RISK = {am.HIGH: "High risk", am.MEDIUM: "Watch", am.LOW: "On track", am.NO_DATA: "No data"}
_DEFINITIONS = [
    ("Attended", "Present plus late classes."),
    ("Classes", "Present, late and absent classes in completed sessions. Excused classes are not counted."),
    ("Attendance %", "Attended divided by classes."),
    ("Risk", "Compared with the subject's attendance target; based on the current percentage and what is still reachable."),
]


def safe_cell(value):
    """Neutralise spreadsheet formula injection: text that starts with = + - @ (or a control char) is prefixed with an apostrophe."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def to_csv(rows: list[dict]) -> bytes:
    """UTF-8 with a BOM so Excel opens non-ASCII names correctly."""
    out = io.StringIO()
    if rows:
        writer = csv.DictWriter(out, fieldnames=list(rows[0]), lineterminator="\n")
        writer.writeheader()
        writer.writerows({k: safe_cell(v) for k, v in r.items()} for r in rows)
    return b"\xef\xbb\xbf" + out.getvalue().encode("utf-8")


def to_xlsx(sheets: dict[str, list[dict]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name[:31])
        if rows:
            ws.append(list(rows[0]))
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for r in rows:
                ws.append([safe_cell(v) for v in r.values()])
            for i, key in enumerate(rows[0], 1):
                ws.column_dimensions[ws.cell(1, i).column_letter].width = min(40, max(len(str(key)), *(len(str(r[key])) for r in rows)) + 2)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def subject_report(teacher_id: int, subject_id: int) -> dict:
    """{subject, generated_at, students, sessions, matrix} - lists of flat dict rows ready for display or export."""
    d = dashboard_service.subject_detail(teacher_id, subject_id)
    ta = analytics_service.teacher_analytics(teacher_id, subject_id, "week")
    students = [{"Student": s["name"], "Attended": s["standing"].attended, "Classes": s["standing"].conducted, "Excused": s["standing"].excused,
                 "Attendance %": s["standing"].percentage, "Target %": s["standing"].target, "Risk": _RISK.get(s["standing"].risk, "No data")} for s in d["students"]]
    summaries = sorted(analytics_service._summaries(subject_id), key=lambda s: (s["started_at"], s["session_id"]))
    sessions = [{"Date": format_local(s["started_at"].isoformat()), "Method": str(s["method"]).replace("_", " ").title(), "Present": s["present"],
                 "Late": s["late"], "Absent": s["absent"], "Excused": s["excused"]} for s in summaries]
    names: dict[str, dict] = {}
    cols: list[str] = []
    for c in sorted(ta.heatmap, key=lambda c: c["order"]):
        names.setdefault(c["student"], {"Student": c["student"]})[c["session"]] = {1: "Present", 0: "Absent", None: "Excused / none"}[c["value"]]
        if c["session"] not in cols:
            cols.append(c["session"])
    matrix = [{"Student": n, **{col: row.get(col, "") for col in cols}} for n, row in sorted(names.items())]
    return {"subject": d["subject"], "generated_at": datetime.now(timezone.utc), "students": students, "sessions": sessions, "matrix": matrix}


def subject_workbook(report: dict) -> bytes:
    sub = report["subject"]
    about = [{"Report": f"{sub['name']} ({sub['subject_code']}) attendance", "Generated": format_local(report["generated_at"].isoformat())}] + \
            [{"Report": k, "Generated": v} for k, v in _DEFINITIONS]
    return to_xlsx({"About": about, "Students": report["students"], "Sessions": report["sessions"], "Attendance matrix": report["matrix"]})


def student_history_rows(student_id: int) -> list[dict]:
    """The student's own records, ready to export (the repository is scoped to the logged-in student by RLS)."""
    from src.ui.widgets import plain_status
    return [{"Date": format_local(r["when"]), "Subject": r["subject"], "Code": r["code"], "Status": plain_status(r["status"])}
            for r in dashboard_service.student_history(student_id)]


def audit_rows(entries: list[dict]) -> list[dict]:
    return [{"When": format_local(a["created_at"]), "Actor": f"{a['actor_role']} {a['actor_id'] or ''}".strip(), "Action": a["action"],
             "Entity": f"{a['entity_type']} {a['entity_id'] or ''}".strip()} for a in entries]
